"""Shared Word styling for the generated documents.

One place decides how a heading, a table or a callout looks, so the demo script
and the walkthrough cannot drift apart typographically.
"""

from __future__ import annotations

import asyncio
import pathlib

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from sqlalchemy import text

from app.core.database import SessionFactory

# The three journeys, named rather than pattern-matched: their identifiers
# follow the book's own sequences, so there is no marker to search for.
CUSTOMERS = ("CUST-CON-207", "CUST-CON-208", "CUST-ENT-036")
ACCOUNTS = ("ACC-10066", "ACC-10067", "ACC-20007")
CASES = ("CASE-000150", "CASE-000151", "CASE-000152")
PTPS = ("PTP-R001253", "PTP-R001254", "PTP-R001255", "PTP-R001256")
PLACEMENTS = ("PL-1092", "PL-1093")
LEGAL = ("LC-2026012", "LC-2026013")
DISPUTES = ("DSP-R001252",)
SESSIONS = ("sess_9c41d0be5f7a4e12ab73", "sess_2f88a51c6d3b47e9c104",
            "sess_71e0b93a4c8d42f5b6a8")
CHAT_FIRST, CHAT_LIVE, CHAT_DANIEL = SESSIONS

INK = RGBColor(0x1B, 0x22, 0x33)
MUTED = RGBColor(0x5B, 0x66, 0x77)
ACCENT = RGBColor(0x2C, 0x5A, 0xD8)
GOOD = RGBColor(0x1B, 0x7F, 0x4B)
WARN = RGBColor(0xB4, 0x54, 0x09)


# --------------------------------------------------------------------------
# Document furniture
# --------------------------------------------------------------------------
def shade(cell, hex_fill: str) -> None:
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:fill"), hex_fill)
    cell._tc.get_or_add_tcPr().append(el)


class Doc:
    def __init__(self) -> None:
        self.doc = Document()
        s = self.doc.sections[0]
        s.top_margin = s.bottom_margin = Inches(0.8)
        s.left_margin = s.right_margin = Inches(0.85)

        normal = self.doc.styles["Normal"]
        normal.font.name = "Calibri"
        normal.font.size = Pt(10.5)
        normal.font.color.rgb = INK
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.line_spacing = 1.12

    # --- blocks ---------------------------------------------------------
    def h(self, txt: str, level: int = 1) -> None:
        p = self.doc.add_heading(txt, level=level)
        for r in p.runs:
            r.font.color.rgb = INK if level > 1 else ACCENT
            r.font.name = "Calibri"
        p.paragraph_format.space_before = Pt(18 if level == 1 else 12)
        p.paragraph_format.space_after = Pt(4)

    def p(self, txt: str, *, italic: bool = False, colour: RGBColor | None = None,
          bold: bool = False, size: float = 10.5) -> None:
        par = self.doc.add_paragraph()
        r = par.add_run(txt)
        r.italic = italic
        r.bold = bold
        r.font.size = Pt(size)
        if colour:
            r.font.color.rgb = colour

    def say(self, txt: str) -> None:
        """A line the presenter can read out."""
        par = self.doc.add_paragraph()
        par.paragraph_format.left_indent = Inches(0.28)
        par.paragraph_format.space_before = Pt(6)
        par.paragraph_format.space_after = Pt(10)
        r = par.add_run("“" + txt + "”")
        r.italic = True
        r.font.size = Pt(10.5)
        r.font.color.rgb = ACCENT

    def bullet(self, txt: str, bold_prefix: str | None = None) -> None:
        par = self.doc.add_paragraph(style="List Bullet")
        par.paragraph_format.space_after = Pt(3)
        if bold_prefix:
            b = par.add_run(bold_prefix)
            b.bold = True
        par.add_run(txt)

    def step(self, n: int, title: str, screen: str) -> None:
        par = self.doc.add_paragraph()
        par.paragraph_format.space_before = Pt(12)
        par.paragraph_format.space_after = Pt(2)
        r = par.add_run(f"{n}.  {title}")
        r.bold = True
        r.font.size = Pt(11.5)
        r.font.color.rgb = INK
        sub = self.doc.add_paragraph()
        sub.paragraph_format.space_after = Pt(6)
        s = sub.add_run(screen)
        s.font.size = Pt(9)
        s.font.color.rgb = MUTED

    def table(self, headers: list[str], rows: list[list[str]],
              widths: list[float] | None = None, right: set[int] | None = None) -> None:
        t = self.doc.add_table(rows=1, cols=len(headers))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.LEFT
        right = right or set()
        for i, head in enumerate(headers):
            c = t.rows[0].cells[i]
            c.text = ""
            run = c.paragraphs[0].add_run(head)
            run.bold = True
            run.font.size = Pt(9)
            run.font.color.rgb = INK
            shade(c, "EEF2F9")
            if i in right:
                c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for row in rows:
            cells = t.add_row().cells
            for i, val in enumerate(row):
                cells[i].text = ""
                run = cells[i].paragraphs[0].add_run(str(val))
                run.font.size = Pt(9.5)
                if i in right:
                    cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        if widths:
            for r_ in t.rows:
                for i, w in enumerate(widths):
                    r_.cells[i].width = Inches(w)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def callout(self, title: str, body: str, colour: RGBColor = WARN) -> None:
        t = self.doc.add_table(rows=1, cols=1)
        t.style = "Table Grid"
        c = t.rows[0].cells[0]
        shade(c, "FFF7EC" if colour is WARN else "EEF6FF")
        c.text = ""
        p1 = c.paragraphs[0]
        r = p1.add_run(title)
        r.bold = True
        r.font.size = Pt(10)
        r.font.color.rgb = colour
        p2 = c.add_paragraph()
        r2 = p2.add_run(body)
        r2.font.size = Pt(9.5)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def code(self, txt: str) -> None:
        par = self.doc.add_paragraph()
        par.paragraph_format.left_indent = Inches(0.2)
        par.paragraph_format.space_before = Pt(4)
        par.paragraph_format.space_after = Pt(10)
        r = par.add_run(txt)
        r.font.name = "Consolas"
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0x24, 0x29, 0x38)

    def rule(self) -> None:
        par = self.doc.add_paragraph()
        par.paragraph_format.space_before = Pt(2)
        par.paragraph_format.space_after = Pt(2)
        pr = par._p.get_or_add_pPr()
        borders = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "6")
        bottom.set(qn("w:color"), "D6DCE6")
        borders.append(bottom)
        pr.append(borders)

    def page_break(self) -> None:
        self.doc.add_page_break()


def money(v) -> str:
    return f"${float(v):,.2f}"


